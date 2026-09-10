from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os


# ============================================================
# CARS24 PROJECT REPORT GENERATOR
# ============================================================

OUTPUT_FILE = "07_Report/Cars24_Used_Car_Analysis_Report.docx"


print("=" * 70)
print("CARS24 PROJECT REPORT GENERATOR")
print("=" * 70)


# ============================================================
# CREATE DOCUMENT
# ============================================================

document = Document()


# ============================================================
# PAGE SETTINGS
# ============================================================

section = document.sections[0]

section.top_margin = Pt(50)
section.bottom_margin = Pt(50)
section.left_margin = Pt(60)
section.right_margin = Pt(60)


# ============================================================
# DEFAULT FONT
# ============================================================

styles = document.styles

styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)


# ============================================================
# TITLE PAGE
# ============================================================

title = document.add_paragraph()

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run(
    "\n\nCARS24 USED CAR\n"
    "MARKET ANALYSIS"
)

run.bold = True
run.font.size = Pt(26)


subtitle = document.add_paragraph()

subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "\nData Analytics & Business Intelligence Project"
)

run.font.size = Pt(14)


details = document.add_paragraph()

details.alignment = WD_ALIGN_PARAGRAPH.CENTER

details.add_run(
    "\n\n\nTools & Technologies\n\n"
    "Python | Selenium | Pandas | NumPy\n"
    "Matplotlib | Seaborn | Power BI\n\n"
    "End-to-End Data Analytics Project"
).font.size = Pt(11)


document.add_page_break()


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def add_heading(text, level=1):

    heading = document.add_heading(
        text,
        level=level
    )

    return heading


def add_paragraph(text):

    paragraph = document.add_paragraph(
        text
    )

    paragraph.paragraph_format.space_after = Pt(6)

    return paragraph


def add_bullet(text):

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.add_run(text)

    return paragraph


# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================

add_heading(
    "1. Executive Summary",
    1
)

add_paragraph(
    "This project analyzes used car listings collected "
    "from Cars24. The project follows a complete data "
    "analytics workflow covering web scraping, raw data "
    "extraction, data cleaning, exploratory data analysis, "
    "data visualization and Power BI dashboard development."
)

add_paragraph(
    "Python was used for data collection, data preparation "
    "and analysis, while Power BI was used to create an "
    "interactive dashboard for business-oriented insights."
)


# ============================================================
# 2. BUSINESS PROBLEM
# ============================================================

add_heading(
    "2. Business Problem",
    1
)

add_paragraph(
    "Used-car marketplaces contain vehicles with different "
    "brands, prices, manufacturing years, fuel types, "
    "transmission types, mileage levels and locations."
)

add_paragraph(
    "The objective of this project is to analyze Cars24 "
    "used-car listings and understand pricing and inventory "
    "patterns within the collected sample."
)


# ============================================================
# 3. OBJECTIVES
# ============================================================

add_heading(
    "3. Project Objectives",
    1
)

objectives = [
    "Collect used-car listing data.",
    "Extract important vehicle attributes.",
    "Clean and prepare the dataset.",
    "Perform exploratory data analysis.",
    "Identify pricing and market patterns.",
    "Create meaningful data visualizations.",
    "Develop an interactive Power BI dashboard.",
    "Generate business-oriented insights."
]

for item in objectives:
    add_bullet(item)


# ============================================================
# 4. DATA COLLECTION
# ============================================================

add_heading(
    "4. Data Collection",
    1
)

add_paragraph(
    "Cars24 used-car listing information was collected using "
    "Python and Selenium. The webpage source was captured "
    "and stored as an HTML file. Listing records were then "
    "extracted from the webpage data."
)

add_paragraph(
    "The raw dataset was stored in "
    "02_Raw_Data/cars24_raw.csv."
)


# ============================================================
# 5. DATA CLEANING
# ============================================================

add_heading(
    "5. Data Cleaning",
    1
)

add_paragraph(
    "The raw dataset was processed using Python and Pandas. "
    "The cleaning workflow prepared the data for analysis "
    "and Power BI."
)

cleaning_steps = [
    "Removed completely empty columns.",
    "Cleaned text fields and unnecessary spaces.",
    "Converted numerical fields to numeric data types.",
    "Checked invalid price values.",
    "Checked invalid mileage values.",
    "Checked vehicle manufacturing years.",
    "Removed duplicate listings.",
    "Reset the dataset index.",
    "Saved the final cleaned dataset."
]

for item in cleaning_steps:
    add_bullet(item)

add_paragraph(
    "The cleaned dataset was stored in "
    "03_Cleaned_Data/cars24_cleaned.csv."
)


# ============================================================
# 6. EXPLORATORY DATA ANALYSIS
# ============================================================

add_heading(
    "6. Exploratory Data Analysis",
    1
)

add_paragraph(
    "Exploratory data analysis was performed to understand "
    "the characteristics of the collected used-car dataset."
)

eda_items = [
    "Brand distribution",
    "Average price by brand",
    "Fuel type distribution",
    "Transmission distribution",
    "Body type distribution",
    "Manufacturing year distribution",
    "Average price by manufacturing year",
    "Mileage statistics",
    "Mileage versus price",
    "Location distribution",
    "City distribution",
    "Ownership distribution",
    "Highest-priced vehicles",
    "Lowest-priced vehicles"
]

for item in eda_items:
    add_bullet(item)


# ============================================================
# 7. VISUALIZATIONS
# ============================================================

add_heading(
    "7. Data Visualizations",
    1
)

add_paragraph(
    "Python was used to create visualizations for the major "
    "analytical questions."
)

visualizations = [
    "Cars by Brand",
    "Average Used Car Price by Brand",
    "Used Car Price Distribution",
    "Fuel Type Distribution",
    "Transmission Distribution",
    "Cars by Manufacturing Year",
    "Average Price by Manufacturing Year",
    "Mileage versus Price"
]

for item in visualizations:
    add_bullet(item)


# ============================================================
# 8. POWER BI DASHBOARD
# ============================================================

add_heading(
    "8. Power BI Dashboard",
    1
)

add_paragraph(
    "An interactive Power BI dashboard was developed using "
    "the cleaned Cars24 dataset."
)

add_heading(
    "Key Performance Indicators",
    2
)

kpis = [
    "Total Cars",
    "Average Car Price",
    "Average Mileage",
    "Average Car Age"
]

for item in kpis:
    add_bullet(item)


add_heading(
    "Dashboard Visuals",
    2
)

dashboard_visuals = [
    "Cars by Brand",
    "Average Price by Brand",
    "Average Price by Manufacturing Year",
    "Fuel Type Distribution",
    "Transmission Distribution",
    "Mileage versus Price"
]

for item in dashboard_visuals:
    add_bullet(item)


add_heading(
    "Interactive Filters",
    2
)

filters = [
    "Brand",
    "Fuel Type",
    "Transmission",
    "Manufacturing Year",
    "City"
]

for item in filters:
    add_bullet(item)


# ============================================================
# 9. KEY INSIGHTS
# ============================================================

add_heading(
    "9. Key Insights",
    1
)

add_paragraph(
    "The analysis provides an overview of brand availability, "
    "used-car pricing, vehicle age, mileage, fuel type, "
    "transmission type and geographic availability."
)

add_paragraph(
    "The detailed numerical findings are based on the "
    "scraped sample dataset used in this project."
)


# ============================================================
# 10. BUSINESS RECOMMENDATIONS
# ============================================================

add_heading(
    "10. Business Recommendations",
    1
)

recommendations = [
    "Monitor inventory availability by vehicle brand.",
    "Compare pricing across different vehicle brands.",
    "Consider vehicle age and mileage when evaluating prices.",
    "Monitor fuel-type and transmission preferences.",
    "Use location-level inventory information for better allocation.",
    "Use interactive dashboards for faster pricing and inventory analysis."
]

for item in recommendations:
    add_bullet(item)


# ============================================================
# 11. TOOLS AND TECHNOLOGIES
# ============================================================

add_heading(
    "11. Tools and Technologies",
    1
)

technologies = [
    "Python",
    "Pandas",
    "NumPy",
    "Selenium",
    "BeautifulSoup",
    "Matplotlib",
    "Seaborn",
    "Microsoft Power BI",
    "Visual Studio Code"
]

for item in technologies:
    add_bullet(item)


# ============================================================
# 12. PROJECT STRUCTURE
# ============================================================

add_heading(
    "12. Project Structure",
    1
)

structure = (
    "Cars24_Used_Car_Analysis/\n\n"
    "├── 01_Web_Scraping/\n"
    "├── 02_Raw_Data/\n"
    "├── 03_Cleaned_Data/\n"
    "├── 04_EDA/\n"
    "├── 05_Visualizations/\n"
    "├── 06_PowerBI/\n"
    "├── 07_Report/\n"
    "├── 08_Presentation/\n"
    "├── main.py\n"
    "├── README.md\n"
    "└── requirements.txt"
)

paragraph = document.add_paragraph()

run = paragraph.add_run(structure)

run.font.name = "Consolas"
run.font.size = Pt(9)


# ============================================================
# 13. LIMITATIONS
# ============================================================

add_heading(
    "13. Limitations",
    1
)

add_paragraph(
    "The analysis is based on a sample of scraped Cars24 "
    "listings. Therefore, the findings should not be "
    "interpreted as a complete representation of the entire "
    "Cars24 inventory or the overall Indian used-car market."
)

add_paragraph(
    "Vehicle listings can change over time as cars are added, "
    "removed or updated."
)


# ============================================================
# 14. CONCLUSION
# ============================================================

add_heading(
    "14. Conclusion",
    1
)

add_paragraph(
    "This project demonstrates an end-to-end data analytics "
    "workflow starting from web data collection and ending "
    "with an interactive business intelligence dashboard."
)

add_paragraph(
    "The project demonstrates practical skills in web "
    "scraping, data extraction, data cleaning, exploratory "
    "data analysis, visualization, Power BI development and "
    "business insight generation."
)


# ============================================================
# SAVE DOCUMENT
# ============================================================

os.makedirs(
    "07_Report",
    exist_ok=True
)

document.save(
    OUTPUT_FILE
)


# ============================================================
# FINAL MESSAGE
# ============================================================

print()
print("=" * 70)
print("REPORT CREATED SUCCESSFULLY")
print("=" * 70)

print()
print(
    f"Report file: {OUTPUT_FILE}"
)

print()
print("DONE")
print("=" * 70)