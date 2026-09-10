from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os


# ============================================================
# CARS24 PROJECT PRESENTATION GENERATOR
# ============================================================

OUTPUT_FILE = (
    "08_Presentation/"
    "Cars24_Used_Car_Analysis_Presentation.pptx"
)


print("=" * 70)
print("CARS24 PROJECT PRESENTATION GENERATOR")
print("=" * 70)


# ============================================================
# CREATE PRESENTATION
# ============================================================

prs = Presentation()

# 16:9 widescreen
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


# ============================================================
# COLORS
# ============================================================

DARK = RGBColor(25, 35, 55)
BLUE = RGBColor(40, 90, 160)
LIGHT = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(90, 100, 115)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def add_background(slide):

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = LIGHT


def add_title(slide, title, subtitle=None):

    # Title
    box = slide.shapes.add_textbox(
        Inches(0.65),
        Inches(0.35),
        Inches(12),
        Inches(0.7)
    )

    text_frame = box.text_frame
    text_frame.clear()

    paragraph = text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(26)
    paragraph.font.bold = True
    paragraph.font.color.rgb = DARK

    if subtitle:

        box2 = slide.shapes.add_textbox(
            Inches(0.68),
            Inches(1.05),
            Inches(11.8),
            Inches(0.45)
        )

        tf = box2.text_frame
        tf.clear()

        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY


def add_footer(slide, slide_number):

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(7.25),
        Inches(13.333),
        Inches(0.25)
    )

    line.fill.solid()
    line.fill.fore_color.rgb = DARK
    line.line.fill.background()

    box = slide.shapes.add_textbox(
        Inches(11.8),
        Inches(7.27),
        Inches(1),
        Inches(0.2)
    )

    p = box.text_frame.paragraphs[0]
    p.text = str(slide_number)
    p.font.size = Pt(9)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.RIGHT


def add_bullets(
    slide,
    items,
    left=0.9,
    top=1.5,
    width=11.5,
    height=5.2
):

    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    tf = box.text_frame
    tf.clear()

    for index, item in enumerate(items):

        if index == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = item
        p.font.size = Pt(18)
        p.font.color.rgb = DARK
        p.space_after = Pt(14)
        p.level = 0


def add_section_box(
    slide,
    title,
    items,
    left,
    top,
    width,
    height
):

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE

    shape.line.color.rgb = BLUE

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(left + 0.25),
        Inches(top + 0.2),
        Inches(width - 0.5),
        Inches(0.45)
    )

    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = BLUE

    # Items
    item_box = slide.shapes.add_textbox(
        Inches(left + 0.3),
        Inches(top + 0.8),
        Inches(width - 0.6),
        Inches(height - 1)
    )

    tf = item_box.text_frame
    tf.clear()

    for index, item in enumerate(items):

        if index == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = "• " + item
        p.font.size = Pt(13)
        p.font.color.rgb = DARK
        p.space_after = Pt(8)


def add_kpi_box(
    slide,
    title,
    value,
    left,
    top,
    width=2.8,
    height=1.5
):

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = BLUE

    title_box = slide.shapes.add_textbox(
        Inches(left + 0.15),
        Inches(top + 0.18),
        Inches(width - 0.3),
        Inches(0.35)
    )

    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GRAY
    p.alignment = PP_ALIGN.CENTER

    value_box = slide.shapes.add_textbox(
        Inches(left + 0.15),
        Inches(top + 0.62),
        Inches(width - 0.3),
        Inches(0.6)
    )

    p = value_box.text_frame.paragraphs[0]
    p.text = value
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK
    p.alignment = PP_ALIGN.CENTER


# ============================================================
# SLIDE 1 — TITLE
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)


# Main title
box = slide.shapes.add_textbox(
    Inches(1),
    Inches(1.7),
    Inches(11.3),
    Inches(1.2)
)

p = box.text_frame.paragraphs[0]

p.text = "CARS24 USED CAR\nMARKET ANALYSIS"

p.font.size = Pt(34)
p.font.bold = True
p.font.color.rgb = DARK
p.alignment = PP_ALIGN.CENTER


box2 = slide.shapes.add_textbox(
    Inches(2),
    Inches(3.4),
    Inches(9.3),
    Inches(0.8)
)

p = box2.text_frame.paragraphs[0]

p.text = (
    "Data Analytics & Business Intelligence Project"
)

p.font.size = Pt(18)
p.font.color.rgb = BLUE
p.alignment = PP_ALIGN.CENTER


box3 = slide.shapes.add_textbox(
    Inches(2),
    Inches(4.5),
    Inches(9.3),
    Inches(1)
)

p = box3.text_frame.paragraphs[0]

p.text = (
    "Python  |  Selenium  |  Pandas  |  NumPy  |  "
    "Matplotlib  |  Power BI"
)

p.font.size = Pt(14)
p.font.color.rgb = GRAY
p.alignment = PP_ALIGN.CENTER


add_footer(slide, 1)


# ============================================================
# SLIDE 2 — PROJECT OVERVIEW
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Project Overview",
    "End-to-end used-car data analytics workflow"
)

add_section_box(
    slide,
    "Project Workflow",
    [
        "Web Scraping",
        "Data Extraction",
        "Data Cleaning",
        "Exploratory Data Analysis",
        "Data Visualization",
        "Power BI Dashboard"
    ],
    0.8,
    1.6,
    5.7,
    4.8
)

add_section_box(
    slide,
    "Project Goal",
    [
        "Understand used-car inventory",
        "Analyze pricing patterns",
        "Study vehicle characteristics",
        "Identify useful business insights"
    ],
    6.8,
    1.6,
    5.7,
    4.8
)

add_footer(slide, 2)


# ============================================================
# SLIDE 3 — BUSINESS PROBLEM
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Business Problem",
    "Understanding used-car inventory and pricing patterns"
)

add_bullets(
    slide,
    [
        "Used-car listings contain multiple attributes that influence customer choice and pricing.",
        "Different brands have different levels of representation in the collected inventory.",
        "Vehicle price can vary according to age, mileage, brand and other characteristics.",
        "Fuel type and transmission provide additional insight into inventory composition.",
        "Location information can help understand where vehicles are available.",
        "The project converts these raw listings into structured, actionable information."
    ],
    top=1.5
)

add_footer(slide, 3)


# ============================================================
# SLIDE 4 — DATA COLLECTION
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Data Collection",
    "Python and Selenium were used to capture Cars24 listing data"
)

add_section_box(
    slide,
    "Technology",
    [
        "Python",
        "Selenium",
        "HTML",
        "Regular Expressions"
    ],
    0.8,
    1.6,
    3.7,
    4.5
)

add_section_box(
    slide,
    "Process",
    [
        "Open Cars24 used-car page",
        "Load listing information",
        "Scroll page to load records",
        "Capture webpage HTML",
        "Extract structured listing records",
        "Save raw CSV dataset"
    ],
    4.8,
    1.6,
    3.7,
    4.5
)

add_section_box(
    slide,
    "Output",
    [
        "cars24_page.html",
        "cars24_raw.csv",
        "20 collected sample listings",
        "Structured vehicle attributes"
    ],
    8.8,
    1.6,
    3.7,
    4.5
)

add_footer(slide, 4)


# ============================================================
# SLIDE 5 — DATA CLEANING
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Data Cleaning",
    "Preparing raw listing data for analysis"
)

add_bullets(
    slide,
    [
        "Removed completely empty columns.",
        "Cleaned text fields and unnecessary spaces.",
        "Converted price, year, mileage and other numerical fields.",
        "Checked invalid price values.",
        "Checked invalid mileage values.",
        "Validated vehicle manufacturing years.",
        "Removed duplicate listings using appointment ID.",
        "Saved the final cleaned dataset for EDA and Power BI."
    ],
    top=1.45
)

add_footer(slide, 5)


# ============================================================
# SLIDE 6 — EDA
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Exploratory Data Analysis",
    "Analyzing the structure and characteristics of the dataset"
)

add_section_box(
    slide,
    "Vehicle Characteristics",
    [
        "Brand",
        "Model",
        "Body Type",
        "Fuel Type",
        "Transmission",
        "Manufacturing Year"
    ],
    0.7,
    1.5,
    3.8,
    4.8
)

add_section_box(
    slide,
    "Pricing & Usage",
    [
        "Average price",
        "Median price",
        "Minimum & maximum price",
        "Mileage statistics",
        "Mileage versus price",
        "Price by manufacturing year"
    ],
    4.75,
    1.5,
    3.8,
    4.8
)

add_section_box(
    slide,
    "Geographic Analysis",
    [
        "Location distribution",
        "City distribution",
        "Inventory availability",
        "Ownership distribution",
        "Highest-priced vehicles",
        "Lowest-priced vehicles"
    ],
    8.8,
    1.5,
    3.8,
    4.8
)

add_footer(slide, 6)


# ============================================================
# SLIDE 7 — VISUALIZATIONS
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Data Visualizations",
    "Python visualizations created to communicate analytical patterns"
)

visuals = [
    "Cars by Brand",
    "Average Price by Brand",
    "Price Distribution",
    "Fuel Type Distribution",
    "Transmission Distribution",
    "Cars by Manufacturing Year",
    "Average Price by Year",
    "Mileage vs Price"
]

add_bullets(
    slide,
    visuals,
    left=1.0,
    top=1.5,
    width=11.2,
    height=5
)

add_footer(slide, 7)


# ============================================================
# SLIDE 8 — POWER BI DASHBOARD
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Power BI Dashboard",
    "Interactive dashboard for used-car market analysis"
)

add_section_box(
    slide,
    "KPIs",
    [
        "Total Cars",
        "Average Car Price",
        "Average Mileage",
        "Average Car Age"
    ],
    0.7,
    1.5,
    3.8,
    4.9
)

add_section_box(
    slide,
    "Visuals",
    [
        "Cars by Brand",
        "Average Price by Brand",
        "Average Price by Year",
        "Fuel Type Distribution",
        "Transmission Distribution",
        "Mileage vs Price"
    ],
    4.75,
    1.5,
    3.8,
    4.9
)

add_section_box(
    slide,
    "Interactive Filters",
    [
        "Brand",
        "Fuel Type",
        "Transmission",
        "Manufacturing Year",
        "City"
    ],
    8.8,
    1.5,
    3.8,
    4.9
)

add_footer(slide, 8)


# ============================================================
# SLIDE 9 — BUSINESS INSIGHTS
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Business Insights",
    "Questions answered through the analysis"
)

add_bullets(
    slide,
    [
        "Which vehicle brands have greater representation in the collected inventory?",
        "How do average used-car prices differ across brands?",
        "What fuel types dominate the collected listings?",
        "What is the distribution between manual and automatic vehicles?",
        "Which manufacturing years are most represented?",
        "How does vehicle mileage relate to listing price?",
        "How is inventory distributed across locations?",
        "How can interactive filtering support faster inventory analysis?"
    ],
    top=1.45
)

add_footer(slide, 9)


# ============================================================
# SLIDE 10 — CONCLUSION & SKILLS
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_background(slide)

add_title(
    slide,
    "Conclusion & Skills",
    "End-to-end practical data analytics project"
)

add_section_box(
    slide,
    "Skills Demonstrated",
    [
        "Web Scraping",
        "Python",
        "Pandas",
        "Data Cleaning",
        "EDA",
        "Data Visualization",
        "Power BI",
        "Business Intelligence"
    ],
    0.8,
    1.5,
    5.7,
    4.9
)

add_section_box(
    slide,
    "Conclusion",
    [
        "Built an end-to-end analytics workflow.",
        "Converted web data into structured information.",
        "Analyzed used-car pricing and inventory characteristics.",
        "Created professional visualizations.",
        "Developed an interactive Power BI dashboard.",
        "Created a foundation for data-driven used-car analysis."
    ],
    6.8,
    1.5,
    5.7,
    4.9
)

add_footer(slide, 10)


# ============================================================
# SAVE
# ============================================================

os.makedirs(
    "08_Presentation",
    exist_ok=True
)

prs.save(
    OUTPUT_FILE
)


# ============================================================
# FINAL MESSAGE
# ============================================================

print()
print("=" * 70)
print("PRESENTATION CREATED SUCCESSFULLY")
print("=" * 70)

print()
print(
    f"Slides created : {len(prs.slides)}"
)

print(
    f"Output file    : {OUTPUT_FILE}"
)

print()
print("=" * 70)
print("DONE")
print("=" * 70)